"""Build current Chapters I-V in the established Word package, without rendering."""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "Finalized Complete Paper WASD - reconciled.docx"
PANDOC = Path(r"C:\Users\JM\AppData\Local\Pandoc\pandoc.exe")
SOURCES = [ROOT / "thesis_docs" / name for name in (
    "Finals_Revised_Paper_WASD.md", "CHAPTER_IV_RESULTS_AND_DISCUSSION.md",
    "CHAPTER_V_SUMMARY_CONCLUSIONS_AND_RECOMMENDATIONS.md")]
RECEIPT = Path(__file__).with_name("current_complete_docx_build_receipt.json")


def sha(payload):
    return hashlib.sha256(payload).hexdigest()


def module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


SYNC = module(ROOT / "thesis_docs/.docx_conversion_qa/sync_final_revised_docx.py", "existing_sync")
BUILDER = module(Path(__file__).with_name("build_reconciled_complete_docx.py"), "existing_tables")


def split_tables(text):
    # Keep manuscript table numbers intact, using explicit parts for wide tables.
    text = re.sub(r"^### Table 4\.(?:1a?|2|3)\..*\n", "", text, flags=re.M)
    lines, output, index = text.splitlines(), [], 0
    while index < len(lines):
        if lines[index].startswith("|") and index+1 < len(lines) and re.match(r"^\|\s*:?-", lines[index+1]):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index]); index += 1
            rows = [BUILDER.parse_table_row(line) for line in block]
            header = rows[0]
            if len(header) == 12 and header[0] == "Method":
                parts = [("Table 4.1 Part A. Ranking Performance and F2", [0,1,2,3,4]),
                         ("Table 4.1 Part B. F2 Operating-Point Performance", [0,5,6,7]),
                         ("Table 4.1 Part C. F2 Confusion Counts", [0,8,9,10,11])]
            elif len(header) == 9 and header[:2] == ["Method", "Policy"]:
                parts = [("Table 4.1a Part A. Policy Cutoffs and Confusion Counts", [0,1,2,3,4,5,6]),
                         ("Table 4.1a Part B. Policy Precision and Recall", [0,1,7,8])]
            elif len(header) == 6 and header[0] == "Comparison":
                parts = [("Table 4.2 Part A. Paired PR-AUC and F2 Differences", [0,1,2,5]),
                         ("Table 4.2 Part B. Paired Precision and Recall Differences", [0,3,4,5])]
            elif len(header) == 9 and header[1] == "Validation PR-AUC":
                parts = [("Table 4.3 Part A. Validation and Test PR-AUC and F2", [0,1,2,3,4]),
                         ("Table 4.3 Part B. Validation and Test Precision and Recall", [0,5,6,7,8])]
            else:
                output.extend(block); continue
            for caption, columns in parts:
                output.extend([f"**{caption}**", "", BUILDER.render_table(rows, columns), ""])
        else:
            output.append(lines[index]); index += 1
    return "\n".join(output)


def widths(header):
    count = len(header)
    if header[:2] == ["Method", "Policy"]:
        return [2600,700,1660,1100,1100,1100,1100] if count == 7 else [4400,1000,1980,1980]
    if count == 2:
        return [6500,2860]
    if count == 3:
        return [2800,3280,3280]
    if count == 4:
        return [2800,2180,2180,2200]
    if count == 5:
        return [2800,1640,1640,1640,1640]
    if count == 7:
        return [2580,900,1000,1140,1140,1140,1460]
    first = 2800
    rest = (9360-first)//(count-1)
    return [first]+[rest]*(count-2)+[9360-first-rest*(count-2)]


def style_table(table):
    SYNC.set_content_table_geometry(table, widths([c.text for c in table.rows[0].cells]))
    pr = table._tbl.tblPr
    borders = pr.find(qn("w:tblBorders"))
    if borders is None:
        borders=OxmlElement("w:tblBorders"); pr.append(borders)
    for edge in ("top","bottom","left","right","insideH","insideV"):
        node=borders.find(qn("w:"+edge))
        if node is None:
            node=OxmlElement("w:"+edge); borders.append(node)
        for name,value in {"val":"single","sz":"4","color":"D9D9D9"}.items():
            node.set(qn("w:"+name),value)
    for row_index,row in enumerate(table.rows):
        row_pr=row._tr.get_or_add_trPr()
        for height in list(row_pr.findall(qn("w:trHeight"))): row_pr.remove(height)
        if row_index==0:
            header=OxmlElement("w:tblHeader"); row_pr.append(header)
        for col,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            BUILDER.set_cell_margins(cell,top=80,start=100,bottom=80,end=100)
            shade=cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            if shade is None:
                shade=OxmlElement("w:shd"); cell._tc.get_or_add_tcPr().append(shade)
            shade.set(qn("w:fill"),"E7E6E6" if row_index==0 else "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT if col==0 else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before=Pt(0)
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.line_spacing=1.0
                for run in paragraph.runs:
                    run.font.size=Pt(9 if len(row.cells)>=4 else 10)
                    run.bold=(row_index==0)


def main():
    original_bytes=OUTPUT.read_bytes()
    source_hashes={str(p.relative_to(ROOT)):sha(p.read_bytes()) for p in SOURCES}
    original=Document(BytesIO(original_bytes))
    original_children=SYNC.direct_children(original)
    first_index,_=SYNC.find_anchor(original,"I. Introduction")
    refs_index,_=SYNC.find_anchor(original,"References")
    cover=original_children[:first_index]
    references=original_children[refs_index:]
    section=deepcopy(original.element.body.sectPr)
    assert len(original.sections)==1
    protected_cover=SYNC.canonical_fingerprint(cover[1:])
    protected_refs=SYNC.canonical_fingerprint(references)
    title=SYNC.markdown_title(SOURCES[0])
    main_md=SOURCES[0].read_text(encoding="utf-8")
    body_md=main_md[main_md.index("# Abstract"):main_md.index("# References")]
    body_md += "\n\n"+SOURCES[1].read_text(encoding="utf-8")+"\n\n"+SOURCES[2].read_text(encoding="utf-8")+"\n\n# References\n"
    body_md=split_tables(body_md)
    body_md=re.sub(r"(?m)^(\d+)\.\s+",r"\1\\. ",body_md)
    body_md=re.sub(r"(?m)^- ","• ",body_md)
    body_md=re.sub(r"(?m)^(\d+\\\. .*)$",r"\1\n",body_md)
    with tempfile.TemporaryDirectory(prefix="groomshield-word-") as temporary:
        temporary=Path(temporary)
        md=temporary/"current_complete.md"
        generated_path=temporary/"generated.docx"
        md.write_text(body_md,encoding="utf-8")
        completed=subprocess.run([str(PANDOC),str(md),"-f",BUILDER.PANDOC_READER,"-t","docx","--reference-doc",str(OUTPUT),"-o",str(generated_path)],capture_output=True,text=True,encoding="utf-8")
        assert completed.returncode==0,completed.stderr
        generated=Document(generated_path)
        end,_=SYNC.find_anchor(generated,"References")
        generated_nodes=SYNC.direct_children(generated)[:end]
        imported=[]
        for node in generated_nodes:
            if node.tag in {SYNC.W_BOOKMARK_START,SYNC.W_BOOKMARK_END}: continue
            assert node.tag in {SYNC.W_P,SYNC.W_TBL}
            assert not any(n.get(SYNC.R_ID) for n in node.iter()),"Unexpected relationship-bound body content"
            cloned=deepcopy(node)
            for tag in (SYNC.W_BOOKMARK_START,SYNC.W_BOOKMARK_END):
                for bookmark in list(cloned.iter(tag)): bookmark.getparent().remove(bookmark)
            imported.append(cloned)
        built=Document(BytesIO(original_bytes))
        body=built.element.body
        for node in list(body): body.remove(node)
        for node in cover: body.append(deepcopy(node))
        SYNC.replace_paragraph_text(body[0],title)
        title_pr=body[0].find(qn("w:pPr"))
        if title_pr is None: title_pr=OxmlElement("w:pPr"); body[0].insert(0,title_pr)
        pstyle=OxmlElement("w:pStyle"); pstyle.set(qn("w:val"),"Title"); title_pr.insert(0,pstyle)
        for node in imported:
            if node.tag==SYNC.W_P:
                p_pr=node.find(qn("w:pPr"))
                p_style=p_pr.find(qn("w:pStyle")) if p_pr is not None else None
                if p_style is not None and p_style.get(qn("w:val"))=="Heading1":
                    page=OxmlElement("w:pageBreakBefore"); p_pr.append(page)
            body.append(node)
        refs_break=OxmlElement("w:p")
        br=OxmlElement("w:br"); br.set(qn("w:type"),"page")
        run=OxmlElement("w:r"); run.append(br); refs_break.append(run); body.append(refs_break)
        for node in references:
            if node.tag!=SYNC.W_SECT_PR: body.append(deepcopy(node))
        body.append(section)
        style_ids={s.get(qn("w:styleId")) for s in built.styles.element}
        for node in generated.styles.element:
            if node.tag==qn("w:style") and node.get(qn("w:styleId")) not in style_ids:
                built.styles.element.append(deepcopy(node)); style_ids.add(node.get(qn("w:styleId")))
        if "Title" not in style_ids:
            style=OxmlElement("w:style"); style.set(qn("w:type"),"paragraph"); style.set(qn("w:styleId"),"Title")
            name=OxmlElement("w:name"); name.set(qn("w:val"),"Title"); style.append(name)
            based=OxmlElement("w:basedOn"); based.set(qn("w:val"),"Normal"); style.append(based)
            built.styles.element.append(style)
        for table in built.tables[1:]: style_table(table)
        staged=temporary/"verified_complete.docx"
        built.save(staged)
        reopened=Document(staged)
        assert reopened.paragraphs[0].text==title
        new_children=SYNC.direct_children(reopened)
        abstract_index,_=SYNC.find_anchor(reopened,"Abstract")
        assert SYNC.canonical_fingerprint(new_children[1:abstract_index])==protected_cover
        new_refs_index,_=SYNC.find_anchor(reopened,"References")
        assert SYNC.canonical_fingerprint(new_children[new_refs_index:])==protected_refs
        assert E.tostring(reopened.element.body.sectPr,method="c14n")==E.tostring(section,method="c14n")
        text=BUILDER.document_text(reopened)
        assert "Primary LSTM" not in text and "Enhanced LSTM" not in text and "**" not in text
        assert not any(token in text for token in ("\\[","\\]","\\(","\\)","\\frac"))
        assert SYNC.count_hyperlinks(reopened)==SYNC.count_hyperlinks(original)
        assert len(reopened.element.findall(".//"+qn("m:oMath")))==len(generated.element.findall(".//"+qn("m:oMath")))
        # Check every generated paragraph (including native equation XML), not only keywords.
        generated_paragraphs=[SYNC.node_text(n) for n in imported if n.tag==SYNC.W_P]
        actual_paragraphs=[SYNC.node_text(n) for n in new_children[abstract_index:new_refs_index] if n.tag==SYNC.W_P and not SYNC.is_page_break(n)]
        assert actual_paragraphs==generated_paragraphs
        generated_tables=[n for n in generated_nodes if n.tag==SYNC.W_TBL]
        assert len(reopened.tables)==len(generated_tables)+1
        for original_table,final_table in zip(generated_tables,reopened.tables[1:]):
            assert [SYNC.node_text(c) for c in original_table.iter(qn("w:tc"))]==[SYNC.node_text(c) for c in final_table._tbl.iter(qn("w:tc"))]
        policy_receipt=json.loads((ROOT/"thesis_docs/.threshold_analysis/test_threshold_policy_comparison.json").read_text())
        policy_tables=[t for t in reopened.tables if len(t.columns)>1 and [c.text for c in t.rows[0].cells][:2]==["Method","Policy"]]
        assert len(policy_tables)==2 and all(len(t.rows)==16 for t in policy_tables)
        for entry in policy_receipt["results"].values():
            for policy,point in entry["policies"].items():
                expected=[entry["name"],policy,f"{point['threshold']:.9f}",str(point['tp']),str(point['fp']),str(point['fn']),f"{point['tn']:,}"]
                expected_rates=[entry["name"],policy,f"{point['precision']:.4f}",f"{point['recall']:.4f}"]
                assert expected in [[c.text for c in row.cells] for row in policy_tables[0].rows]
                assert expected_rates in [[c.text for c in row.cells] for row in policy_tables[1].rows]
        # Transfer only the authored XML parts into the unchanged original package.
        final=BytesIO()
        changed_parts={"word/document.xml","word/styles.xml"}
        with ZipFile(BytesIO(original_bytes)) as source_package, ZipFile(staged) as staged_package, ZipFile(final,"w") as result:
            for info in source_package.infolist():
                result.writestr(info,staged_package.read(info.filename) if info.filename in changed_parts else source_package.read(info.filename))
        final_bytes=final.getvalue()
        with ZipFile(BytesIO(final_bytes)) as result,ZipFile(BytesIO(original_bytes)) as reference:
            assert result.testzip() is None
            assert all(result.read(name)==reference.read(name) for name in reference.namelist() if name not in changed_parts)
        assert sha(OUTPUT.read_bytes())==sha(original_bytes),"Reference changed during build"
        assert all(sha(p.read_bytes())==source_hashes[str(p.relative_to(ROOT))] for p in SOURCES)
        OUTPUT.write_bytes(final_bytes)
        receipt=dict(created_utc=datetime.now(timezone.utc).isoformat(),output=str(OUTPUT),output_sha256=sha(final_bytes),previous_output_sha256=sha(original_bytes),sources=source_hashes,
                     verified=True,visual_rendering=False,chapters=5,abstract_included=True,policy_rows_verified=15,
                     cover_authors_preserved=True,references_preserved=True,unchanged_package_parts_preserved=True,
                     edited_package_parts=sorted(changed_parts),native_equations=len(reopened.element.findall(".//"+qn("m:oMath"))),
                     table_shapes=[[len(t.rows),len(t.columns)] for t in reopened.tables],pandoc_warnings=completed.stderr,
                     model_training=False,model_inference=False,active_policy_changed=False)
        RECEIPT.write_text(json.dumps(receipt,indent=2)+"\n",encoding="utf-8")
        print(json.dumps(receipt,indent=2))


if __name__=="__main__": main()
