"""Build the one-page GroomShield abstract submission from ABSTRACT_F2.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ABSTRACT_F2.md"
OUTPUT = ROOT / "docx" / "GroomShield_Abstract_Submission.docx"
ICON_9 = ROOT / "assets" / "UN_SDG_9_English.png"
ICON_16 = ROOT / "assets" / "UN_SDG_16_English.png"


def field(text: str, label: str, next_label: str | None = None) -> str:
    start = re.escape(f"**{label}:**")
    end = re.escape(f"**{next_label}:**") if next_label else r"\Z"
    match = re.search(start + r"\s*(.*?)\s*(?=" + end + r")", text, re.S)
    if match is None:
        raise ValueError(f"Missing field: {label}")
    return match.group(1).strip()


def set_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    font_spec = run._element.get_or_add_rPr().rFonts
    if font_spec is not None:
        font_spec.set(qn("w:ascii"), "Arial")
        font_spec.set(qn("w:hAnsi"), "Arial")


def add_label(doc: Document, label: str, before: float = 7) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(label), 10, bold=True)


def add_name_lines(doc: Document, names: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    for index, name in enumerate(names):
        if index:
            paragraph.add_run().add_break()
        set_font(paragraph.add_run(name), 9.5)


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    title = field(text, "Research Title", "SDG Alignment")
    author_lines = field(text, "Authors", "Co-Authors")
    coauthor_lines = field(text, "Co-Authors", "Research Abstract")
    authors = [part.strip() for part in author_lines.replace("<br>", "\n").splitlines() if part.strip()]
    coauthors = [part.strip() for part in coauthor_lines.replace("<br>", "\n").splitlines() if part.strip()]
    abstract = " ".join(field(text, "Research Abstract", "Keywords").split())
    keywords = field(text, "Keywords")

    if len(authors) != 4 or coauthors != ["Manuel L. Calimlim Jr.", "Lorena Rabago"]:
        raise ValueError("Author or co-author block does not match the confirmed submission list")
    if not ICON_9.is_file() or not ICON_16.is_file():
        raise FileNotFoundError("Official SDG icon files are required")
    if OUTPUT.exists():
        raise FileExistsError(f"Output already exists: {OUTPUT}")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.10
    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(10.5)
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run("Research Title: "), 10.5, bold=True)
    set_font(p.add_run(title), 10.5)

    icons = doc.add_paragraph()
    icons.paragraph_format.space_before = Pt(5)
    icons.paragraph_format.space_after = Pt(10)
    icons.paragraph_format.keep_with_next = True
    for index, (path, description) in enumerate(
        [
            (ICON_9, "United Nations Sustainable Development Goal 9 icon"),
            (ICON_16, "United Nations Sustainable Development Goal 16 icon"),
        ]
    ):
        if index:
            icons.add_run("  ")
        shape = icons.add_run().add_picture(str(path), width=Inches(0.67))
        shape._inline.docPr.set("descr", description)

    team = doc.add_paragraph()
    team.paragraph_format.space_after = Pt(3)
    set_font(team.add_run("Team WASD"), 10, bold=True)

    add_label(doc, "Authors:", before=0)
    add_name_lines(doc, authors)
    add_label(doc, "Co-Authors:", before=8)
    add_name_lines(doc, coauthors)

    add_label(doc, "Research Abstract:", before=11)
    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.16
    body.paragraph_format.space_after = Pt(0)
    set_font(body.add_run(abstract), 10)

    key_para = doc.add_paragraph()
    key_para.paragraph_format.space_before = Pt(10)
    key_para.paragraph_format.space_after = Pt(0)
    set_font(key_para.add_run("Keywords: "), 9.5, bold=True)
    set_font(key_para.add_run(keywords), 9.5)

    # The UN requests a site link and non-endorsement statement when its icons
    # appear in online or printed material.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.line_spacing = 1.0
    set_font(footer.add_run("SDG icons: United Nations, https://www.un.org/sustainabledevelopment. "), 7.2)
    set_font(
        footer.add_run(
            "The content of this publication has not been approved by the United Nations "
            "and does not reflect the views of the United Nations or its officials or Member States."
        ),
        7.2,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
